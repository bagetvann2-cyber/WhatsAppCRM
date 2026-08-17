# -*- coding: utf-8 -*-
"""Конвертация ТЗ из markdown в docx с оформлением."""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

SRC = r"C:\Users\BagetPC_2\Desktop\WhatsAppCRM\docs\TZ-WhatsApp-CRM.md"
DST = r"C:\Users\BagetPC_2\Desktop\WhatsAppCRM\docs\TZ-WhatsApp-CRM.docx"

ACCENT = RGBColor(0x0F, 0x6B, 0x5C)
INK = RGBColor(0x1A, 0x1A, 0x1A)
MUTED = RGBColor(0x66, 0x66, 0x66)
HEADER_FILL = "E2F0EC"


def shade(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)


def add_runs(paragraph, text, base_size=10.5, bold_all=False):
    """Разбирает **bold**, `code`, *italic* и добавляет runs."""
    token = re.compile(r"(\*\*.+?\*\*|`.+?`|\*[^*]+?\*)")
    for part in token.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(base_size - 1)
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1])
            r.italic = True
        else:
            r = paragraph.add_run(part)
        r.font.size = Pt(base_size)
        if bold_all:
            r.bold = True


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Segoe UI"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = INK
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Segoe UI")
    pf = normal.paragraph_format
    pf.space_after = Pt(6)
    pf.line_spacing = 1.15

    specs = [
        ("Heading 1", 20, True, ACCENT, 18, 8),
        ("Heading 2", 15, True, ACCENT, 16, 6),
        ("Heading 3", 12, True, INK, 12, 4),
    ]
    for name, size, bold, color, before, after in specs:
        st = doc.styles[name]
        st.font.name = "Segoe UI Semibold" if bold else "Segoe UI"
        st.font.size = Pt(size)
        st.font.bold = bold
        st.font.color.rgb = color
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def flush_table(doc, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = table.rows[0].cells
    for i, text in enumerate(header):
        hdr[i].text = ""
        p = hdr[i].paragraphs[0]
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, text, base_size=9.5, bold_all=True)
        shade(hdr[i], HEADER_FILL)
    for row in body:
        cells = table.add_row().cells
        for i, text in enumerate(row[: len(header)]):
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(2)
            add_runs(p, text, base_size=9.5)
    doc.add_paragraph()


def main():
    # Без аргументов собирается ТЗ. С аргументами: md2docx.py <исходник.md> [результат.docx]
    src = sys.argv[1] if len(sys.argv) > 1 else SRC
    dst = sys.argv[2] if len(sys.argv) > 2 else (
        DST if src == SRC else re.sub(r"\.md$", ".docx", src)
    )

    with open(src, encoding="utf-8") as fh:
        lines = fh.read().splitlines()

    doc = Document()
    for s in doc.sections:
        s.top_margin = Cm(2)
        s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2)
        s.right_margin = Cm(2.2)
    setup_styles(doc)

    table_buf = []
    in_table = False

    for raw in lines:
        line = raw.rstrip()
        is_table_row = line.startswith("|") and line.endswith("|")

        if in_table and not is_table_row:
            flush_table(doc, table_buf)
            table_buf, in_table = [], False

        if is_table_row:
            cells = [c.strip() for c in line.strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                continue
            table_buf.append(cells)
            in_table = True
            continue

        if not line.strip():
            continue

        if line.startswith("---"):
            continue

        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
        elif line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
        elif line.startswith("# "):
            h = doc.add_heading("", level=1)
            add_runs(h, line[2:].strip(), base_size=20, bold_all=True)
            for r in h.runs:
                r.font.color.rgb = ACCENT
        elif re.match(r"^\d+\.\s", line):
            p = doc.add_paragraph(style="List Number")
            add_runs(p, re.sub(r"^\d+\.\s", "", line))
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style="List Bullet")
            add_runs(p, line[2:])
        elif line.startswith("*") and line.endswith("*") and not line.startswith("**"):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, line)
            for r in p.runs:
                r.font.color.rgb = MUTED
                r.italic = True
        else:
            p = doc.add_paragraph()
            add_runs(p, line)

    if table_buf:
        flush_table(doc, table_buf)

    doc.save(dst)
    print("saved:", dst)


if __name__ == "__main__":
    sys.exit(main())
